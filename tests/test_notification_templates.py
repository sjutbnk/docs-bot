from pathlib import Path

from docx import Document

import docx_populator
import generator


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = PROJECT_ROOT / "bot" / "templates"


def _test_data(employee_initiative=False):
    return {
        "full_name": "Иванов Иван Иванович",
        "phone": "89001234567",
        "citizenship": "Узбекистан",
        "birth_date": "01.01.1990",
        "passport_series": "FA",
        "passport_number": "123456789",
        "passport_issue_date": "01.02.2021",
        "passport_issued_by": "УМВД РОССИИ",
        "patent_series": "30",
        "patent_number": "1234567890",
        "patent_issue_date": "01.02.2026",
        "patent_expiry_date": "01.02.2027",
        "profession": "Овощевод",
        "contract_date": "03.06.2026",
        "contract_end_date": "30.06.2026",
        "termination_employee_initiative": employee_initiative,
        "employer_type": "ИП",
        "employer_name": "ИП Тестов Андрей Андреевич",
        "employer_inn": "123456789012",
        "employer_ogrn": "123456789012345",
        "employer_address": "Астраханская область, тестовый адрес",
        "employer_passport_series": "1219",
        "employer_passport_number": "810165",
        "employer_passport_issue_date": "26.03.2020",
        "employer_passport_issued_by": "УМВД России",
        "work_address": "Астраханская область, тестовый адрес",
    }


def _document_text(doc):
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc)
                    parts.append(cell.text)
    return "\n".join(parts).upper()


def test_new_conclusion_form_is_filled_with_current_data():
    doc = Document(TEMPLATES / "template_conclusion_new.docx")

    docx_populator.fill_new_conclusion_document(doc, _test_data())
    text = _document_text(doc)

    assert "".join(
        cell.text for cell in docx_populator._unique_cells(doc.tables[22].rows[0])[1:7]
    ) == "ИВАНОВ"
    assert "ТЕСТОВ" in text
    assert "БЛОХИН" not in text
    assert "АЛИМОВ" not in text
    assert doc.tables[46].rows[1].cells[1].text == "0"
    assert doc.tables[46].rows[1].cells[2].text == "3"


def test_new_termination_form_replaces_sample_and_sets_initiative():
    doc = Document(TEMPLATES / "template_termination_new.docx")

    docx_populator.fill_new_termination_document(doc, _test_data(employee_initiative=True))
    text = _document_text(doc)

    assert "".join(
        cell.text for cell in docx_populator._unique_cells(doc.tables[22].rows[0])[1:7]
    ) == "ИВАНОВ"
    assert "ТЕСТОВ" in text
    assert "БЛОХИН" not in text
    assert "АЛИМОВ" not in text
    reason_cells = docx_populator._unique_cells(doc.tables[48].rows[0])
    assert reason_cells[1].text == "X"
    assert reason_cells[3].text == ""
    assert doc.tables[46].rows[1].cells[1].text == "3"
    assert doc.tables[46].rows[1].cells[2].text == "0"


def test_new_termination_form_marks_no_when_not_employee_initiative():
    doc = Document(TEMPLATES / "template_termination_new.docx")

    docx_populator.fill_new_termination_document(doc, _test_data(employee_initiative=False))

    reason_cells = docx_populator._unique_cells(doc.tables[48].rows[0])
    assert reason_cells[1].text == ""
    assert reason_cells[3].text == "X"


def test_generator_uses_revised_forms_and_removes_sample_data(tmp_path):
    paths = generator.generate_documents(_test_data(), str(tmp_path))

    assert len(paths) == 4
    assert all(Path(path).is_file() and Path(path).stat().st_size > 0 for path in paths)

    for path in paths[1:3]:
        doc = Document(path)
        text = _document_text(doc)
        surname = "".join(
            cell.text for cell in docx_populator._unique_cells(doc.tables[22].rows[0])[1:7]
        )
        assert surname == "ИВАНОВ"
        assert "ТЕСТОВ" in text
        assert "БЛОХИН" not in text
        assert "АЛИМОВ" not in text
