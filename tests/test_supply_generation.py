"""Offline smoke tests for document generation."""

import zipfile
from pathlib import Path

from docx import Document

import supply_generator
import utils


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _party(prefix: str) -> dict:
    return {
        f"{prefix}_type": "ИП",
        f"{prefix}_name": f"ИП Иванов Иван Иванович",
        f"{prefix}_inn": "123456789012",
        f"{prefix}_ogrn": "123456789012345",
        f"{prefix}_address": "Астраханская область, г. Ахтубинск",
        f"{prefix}_rs": "40802810905000025734",
        f"{prefix}_ks": "30101810500000000602",
        f"{prefix}_bik": "041203602",
        f"{prefix}_bank": "Тестовый банк",
        f"{prefix}_passport_series": "1219",
        f"{prefix}_passport_number": "810165",
        f"{prefix}_passport_issued_by": "УМВД России",
        f"{prefix}_passport_issue_date": "26.03.2020",
    }


def _contract_data(contract_type: str) -> dict:
    data = {
        **_party("supplier"),
        **_party("buyer"),
        "contract_type": contract_type,
        "contract_start_date": "01.06.2026",
        "contract_end_date": "30.06.2026",
    }
    data["buyer_name"] = "ИП Петров Петр Петрович"
    return data


def test_supply_contract_templates_generate_documents(tmp_path, monkeypatch):
    monkeypatch.setattr(supply_generator.config, "TEMPLATES_DIR", str(PROJECT_ROOT / "bot" / "templates"))

    for contract_type in ("vat", "no_vat"):
        output_path = Path(
            supply_generator.generate_supply_contract(
                _contract_data(contract_type), str(tmp_path / contract_type)
            )
        )

        assert output_path.is_file()
        assert output_path.stat().st_size > 0
        with zipfile.ZipFile(output_path) as document:
            assert "[Content_Types].xml" in document.namelist()
            assert "word/document.xml" in document.namelist()
            xml = document.read("word/document.xml").decode("utf-8")
            assert "{{" not in xml
            assert "supplier_party_intro" not in xml


def test_supply_dates_require_real_ordered_calendar_dates():
    assert utils.is_valid_date("29.02.2028")
    assert not utils.is_valid_date("29.02.2027")
    assert utils.dates_are_ordered("01.06.2026", "30.06.2026")
    assert not utils.dates_are_ordered("30.06.2026", "01.06.2026")


def test_farmer_signature_has_no_position_wording():
    assert supply_generator._ip_signature("Иванов Иван Иванович") == "ИП Иванов И.И."


def test_customer_supply_form_uses_current_parties_and_requisites(tmp_path):
    data = _contract_data("no_vat")
    data["supplier_name"] = "ИП глава КФХ Иванов Иван Иванович"
    data["buyer_type"] = "ЮрЛицо"
    data["buyer_name"] = 'ООО "Тестовая компания"'
    data["buyer_director"] = "Петров Петр Петрович"

    output_path = supply_generator.generate_supply_contract(data, str(tmp_path))
    doc = Document(output_path)
    text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )

    assert doc.paragraphs[0].text == "ДОГОВОР ПОСТАВКИ"
    assert "глава крестьянского (фермерского) хозяйства Иванов Иван Иванович" in text
    assert "ИП Иванов И.И." in text
    assert "Тестовая компания" in text
    assert "именуемое в дальнейшем «Покупатель»" in doc.paragraphs[4].text
    assert "123456789012345" in text
    assert "30" in doc.paragraphs[56].text
    assert "без НДС" in text
    assert doc.tables[0].cell(2, 0).text.startswith("ИП Иванов И.И.")
    assert "Глава" not in doc.tables[0].cell(2, 0).text
    for sample_value in ("Ким Владимир", "Фуд-Сервис", "Зейналова", "001466539"):
        assert sample_value not in text
    assert "{{" not in text
